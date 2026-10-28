"""Clicks "Download weekly review (PPT)" and checks the deck against an independent count from the data.
Run: python tests/report_check.py <repo path> [output.pptx]"""
import json, subprocess, sys, time
from pptx import Presentation
from playwright.sync_api import sync_playwright

ROOT = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "/tmp/sea_weekly_review_check.pptx"
D = json.load(open(ROOT + "/data/details.json"))
from datetime import datetime, timedelta
asof = (datetime.strptime(D["generated_at"], "%Y-%m-%dT%H:%M:%SZ") + timedelta(hours=7)).strftime("%Y-%m-%d")  # Bangkok date
y, m = int(asof[:4]), int(asof[5:7])
fy0 = y if m >= 4 else y - 1
c0, c1 = f"{fy0}-04-01", f"{fy0 + 1}-04-01"
active = [r for r in D["active"] if r["stage"] != "On-Hold"]
won = [r for r in D["closures"] if c0 <= r["close_date"] < c1]
props = [r for r in active if r["stage"] == "Proposal Sent"]
val = lambda r: r["year1_total"] if r.get("year1_total") is not None else (r.get("amount") or 0)

fails = []
srv = subprocess.Popen([sys.executable, "-m", "http.server", "8772"], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
time.sleep(1.5)
try:
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(accept_downloads=True)
        errs = []; pg.on("pageerror", lambda e: errs.append(str(e)))
        pg.goto("http://localhost:8772/index.html"); pg.wait_for_timeout(800)
        with pg.expect_download(timeout=60000) as dl:
            pg.click("#wsr-btn")
        dl.value.save_as(OUT)
        if errs: fails.append(f"page errors: {errs}")
        b.close()
finally:
    srv.terminate()

prs = Presentation(OUT)
texts = []
for s in prs.slides:
    t = []
    for sh in s.shapes:
        if sh.has_text_frame: t.append(sh.text_frame.text)
        if sh.has_table: t += [c.text for row in sh.table.rows for c in row.cells]
    texts.append("\n".join(t))
alltext = "\n".join(texts)
print(f"slides: {len(prs.slides)}")
checks = [(f"{len(won)} contracts signed in FY{str(fy0 + 1)[2:]}", "contracts title"),
          (f"{len(props)} proposals out", "proposals title"),
          (f"Open pipeline: {len(active)} deals", "pipeline title"),
          (f"{round(sum(val(r) for r in active)):,}", "pipeline total"),
          (f"{round(sum(val(r) for r in won)):,}", "contracts total")]
for needle, what in checks:
    if needle not in alltext: fails.append(f"{what}: '{needle}' not found")
for r in props:
    if (r.get("company") or r["deal"]) not in alltext: fails.append(f"proposal missing: {r['deal']}")
if "On-Hold" in alltext.replace("On-Hold deals excluded", "").replace("(On-Hold excluded)", ""): fails.append("an On-Hold deal appears in the deck")
print(f"Failures: {len(fails)}")
for f in fails: print("  FAIL", f)
